"""Results Approved: ชื่อไม่ซ้ำ, มีรอบ QF, เที่ยวพิเศษ, ถอนตัว, อันดับสุดท้าย. Run: python -m unittest tests.test_results_approved -v"""
import os, re, tempfile, unittest
from datetime import date, datetime

_test_dir = tempfile.TemporaryDirectory(prefix="shooting-ra-tests-")
os.environ.setdefault("DATABASE_URL", "sqlite:///" + _test_dir.name + "/test.db")
os.environ.setdefault("SECRET_KEY", "test-only-secret")

from flask import g
from app import (app, db, User, Event, Athlete, ScoreEntry, ScoreSignature, BracketMatch, TieBreakEntry,
                 ensure_round_entries, build_results_approved_context, make_results_approved_docx)


class ResultsApprovedTests(unittest.TestCase):
    def setUp(self):
        app.config.update(TESTING=True)
        self.ctx = app.app_context(); self.ctx.push()
        with db.engine.connect() as conn:
            conn.exec_driver_sql("PRAGMA foreign_keys=OFF")
            db.metadata.drop_all(bind=conn); db.metadata.create_all(bind=conn); conn.commit()
        self.admin = User(username="adm", role="admin", password_hash="!UNSET!"); db.session.add(self.admin)
        self.e = Event(name="คำมอกหลวงเกมส์ · กีฬาเปตอง Shooting ม.ปลาย ชาย", event_group="ม.ปลาย", category="ชาย",
                       competition_date=date(2026, 10, 24), location="สาธิตมหาวิทยาลัยพะเยา", lane_count=6,
                       direct_qualifiers=8, has_round_two=False)
        db.session.add(self.e); db.session.flush()
        self.ath = []
        scores = [32, 27, 24, 21, 20, 19, 16, 15, 11, 11, 10, 9, 8, 0]
        for i, sc in enumerate(scores, start=1):
            wd = i == len(scores)
            a = Athlete(event_id=self.e.id, bib_no=str(i), name=("SCHOOL%d" % i) if wd else f"นาย A{i}",
                        affiliation=f"SCHOOL{i}", start_order=i, lane_no=(i - 1) % 6 + 1, lane_order=1, withdrawn=wd)
            db.session.add(a); db.session.flush()
            self.set_round(a.id, 1, sc)
            self.ath.append(a)
        db.session.commit()

    def set_round(self, aid, rnd, total):
        ensure_round_entries(aid, rnd)
        left = total
        for en in ScoreEntry.query.filter_by(athlete_id=aid, round_no=rnd).order_by(ScoreEntry.id):
            v = 5 if left >= 5 else (3 if left >= 3 else (1 if left >= 1 else 0))
            en.score, en.is_scored = v, True
            left -= v
        db.session.add(ScoreSignature(athlete_id=aid, round_no=rnd, started_at=datetime.utcnow(), finished_at=datetime.utcnow(), bypass_signed=True))

    def tearDown(self):
        db.session.remove(); self.ctx.pop()

    def build_bracket(self):
        A = [a.id for a in self.ath]
        qf = [(A[0], A[7], 11, 21), (A[3], A[4], 12, 9), (A[2], A[5], 13, 10), (A[1], A[6], 11, 11)]
        winners = []
        for no, (x, y, sx, sy) in enumerate(qf, start=1):
            self.set_round(x, 3, sx); self.set_round(y, 3, sy)
            w = x if sx > sy else y
            if no == 4:
                for st in range(1, 6):
                    db.session.add(TieBreakEntry(athlete_id=x, round_no=3, station_no=st, score=3 if st == 1 else 0))
                    db.session.add(TieBreakEntry(athlete_id=y, round_no=3, station_no=st, score=0))
                w = x
            winners.append(w)
            db.session.add(BracketMatch(event_id=self.e.id, round_name="QF", match_no=no, athlete_a_id=x, athlete_b_id=y, winner_id=w))
        sf = [(winners[0], winners[1], 10, 11), (winners[2], winners[3], 9, 23)]
        fin = []
        for no, (x, y, sx, sy) in enumerate(sf, start=1):
            self.set_round(x, 4, sx); self.set_round(y, 4, sy)
            fin.append(y)
            db.session.add(BracketMatch(event_id=self.e.id, round_name="SF", match_no=no, athlete_a_id=x, athlete_b_id=y, winner_id=y))
        self.set_round(fin[0], 5, 17); self.set_round(fin[1], 5, 15)
        db.session.add(BracketMatch(event_id=self.e.id, round_name="F", match_no=1, athlete_a_id=fin[0], athlete_b_id=fin[1], winner_id=fin[0]))
        db.session.commit()

    def client(self):
        g.pop("_login_user", None); g.pop("shooting_request_cache", None)
        c = app.test_client()
        with c.session_transaction() as s:
            s["_user_id"] = str(self.admin.id); s["_fresh"] = True
        return c

    def test_report(self):
        self.build_bracket()
        with app.test_request_context():
            ctx = build_results_approved_context(self.e)
        self.assertEqual(ctx["event_title"], "คำมอกหลวงเกมส์ · กีฬาเปตอง SHOOTING ม.ปลาย ชาย")
        self.assertEqual(ctx["cover_title"], "คำมอกหลวงเกมส์")
        self.assertEqual([t for t, _ in ctx["ko_sections"]], ["QUARTERFINAL ROUND", "SEMIFINAL ROUND", "FINAL ROUND"])
        qf4 = [r for r in ctx["bracket_rows"] if r["round_name"] == "QF" and r["match_no"] == 4][0]
        self.assertEqual(qf4["points_a_text"], "11 (7M 3)")
        self.assertEqual(qf4["points_b_text"], "11 (7M 0)")
        wd = ctx["qf1_rows"][-1]
        self.assertEqual((wd["points"], wd["rank"], wd["name"]), ("WD", "WD", ""))
        ranks = [(r["rank"], r["country"]) for r in ctx["final_ranking_rows"]]
        self.assertEqual(ranks[0], (1, "SCHOOL4"))
        self.assertEqual(ranks[1], (2, "SCHOOL2"))
        self.assertEqual({r for r, _ in ranks[2:4]}, {3})
        self.assertEqual({r for r, _ in ranks[4:8]}, {5})
        self.assertEqual(ranks[8][0], 9); self.assertEqual(ranks[9][0], 9)  # คัดเลือกเท่ากัน = อันดับร่วม
        self.assertEqual(ranks[-1], ("WD", "SCHOOL14"))

        html = self.client().get(f"/events/{self.e.id}/results-approved").get_data(as_text=True)
        self.assertNotIn("SEA GAMES", html)
        self.assertNotIn("ม.ปลาย ชาย คำมอกหลวง", html)
        self.assertIn("QUARTERFINAL ROUND", html)
        self.assertIn("11 (7M 3)", html)
        self.assertIn("FINAL RANKING", html)
        self.assertNotIn("UMPIRE", html)  # ยังไม่กรอกกรรมการ ไม่พิมพ์หน้าว่าง
        r = self.client().get(f"/events/{self.e.id}/results-approved.docx")
        self.assertEqual(r.status_code, 200)
        from docx import Document
        from io import BytesIO
        doc = Document(BytesIO(r.data))
        text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        for needle in ("QUARTERFINAL ROUND", "11 (7M 3)", "FINAL RANKING", "คำมอกหลวงเกมส์", "WD"):
            self.assertIn(needle, text)


if __name__ == "__main__":
    unittest.main()
