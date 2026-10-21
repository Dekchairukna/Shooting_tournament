import os
import re
import tempfile
import unittest
from datetime import date
from io import BytesIO

from openpyxl import load_workbook

_test_dir = tempfile.TemporaryDirectory(prefix="shooting-bulk-settings-")
os.environ["DATABASE_URL"] = "sqlite:///" + _test_dir.name + "/test.db"
os.environ.setdefault("SECRET_KEY", "test-only-secret")

from app import app, db, User, Event, Athlete, BracketMatch, SiteTheme


class BulkEventSettingsTests(unittest.TestCase):
    def setUp(self):
        app.config.update(TESTING=True)
        self.ctx = app.app_context()
        self.ctx.push()
        db.drop_all()
        db.create_all()
        user = User(username="boss", role="superadmin", password_hash="!UNSET!")
        db.session.add(user)
        self.events = []
        for index in range(1, 4):
            event = Event(name=f"รายการ {index}", event_group="ทั่วไป", category="ชาย",
                          competition_date=date(2026, 10, 20), location="",
                          lane_count=2, direct_qualifiers=index, has_round_two=True,
                          round_two_cutoff_rank=16, round_two_mode="cutoff",
                          next_round_label="รอบ 8 คน", round_two_advancers=4,
                          overview_sort_by_rank=True, overview_show_cut_lines=True)
            db.session.add(event)
            db.session.flush()
            db.session.add(Athlete(event_id=event.id, bib_no="1", name=f"ทีม {index}",
                                   affiliation="ทดสอบ", start_order=1, lane_no=1, lane_order=1))
            self.events.append(event)
        db.session.add(BracketMatch(event_id=self.events[0].id, round_name="QF", match_no=1))
        db.session.commit()
        self.client = app.test_client()
        with self.client.session_transaction() as session:
            session["_user_id"] = str(user.id)
            session["_fresh"] = True
        page = self.client.get("/").get_data(as_text=True)
        token = re.search(r'name="csrf-token" content="([^"]+)"', page).group(1)
        self.client.environ_base["HTTP_X_CSRF_TOKEN"] = token

    def tearDown(self):
        db.session.remove()
        self.ctx.pop()

    def test_index_has_multi_select_actions(self):
        page = self.client.get("/").get_data(as_text=True)
        self.assertIn("event-select-all", page)
        self.assertIn("ตั้งค่าร่วมกัน", page)
        self.assertIn("เปิดดูรวมกัน", page)

    def test_only_checked_settings_apply_to_selected_events(self):
        ids = f"{self.events[0].id},{self.events[1].id}"
        response = self.client.post("/events/bulk-settings", data={
            "ids": ids,
            "apply_direct_qualifiers": "on",
            "direct_qualifiers": "8",
            "apply_overview_sort_by_rank": "on",
            "overview_sort_by_rank": "no",
            "has_round_two": "no",
        })
        self.assertEqual(response.status_code, 302)
        db.session.expire_all()
        first, second, third = [db.session.get(Event, event.id) for event in self.events]
        self.assertEqual((first.direct_qualifiers, second.direct_qualifiers), (8, 8))
        self.assertEqual(third.direct_qualifiers, 3)
        self.assertFalse(first.overview_sort_by_rank)
        self.assertFalse(second.overview_sort_by_rank)
        self.assertTrue(third.overview_sort_by_rank)
        self.assertTrue(first.has_round_two)  # ไม่ติ๊กหัวข้อนี้ จึงไม่เปลี่ยน
        self.assertEqual(BracketMatch.query.filter_by(event_id=first.id).count(), 0)

    def test_combined_overview_contains_only_selected_events(self):
        url = f"/events/combined-overview?ids={self.events[1].id},{self.events[0].id}"
        page = self.client.get(url).get_data(as_text=True)
        self.assertIn("รายการ 1", page)
        self.assertIn("รายการ 2", page)
        self.assertNotIn("รายการ 3", page)
        self.assertIn("ทีม 1", page)

    def test_overview_results_export_is_valid_excel(self):
        event = self.events[0]
        response = self.client.get(f"/events/{event.id}/overview.xlsx?round=1")
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response.content_type)
        workbook = load_workbook(BytesIO(response.data), data_only=False)
        self.assertEqual(workbook.sheetnames, ["TIR QUALIFICATION", "TIR REPECHAGE", "TIR FINAL "])
        sheet = workbook["TIR QUALIFICATION"]
        self.assertIn(event.name, sheet["E2"].value)
        self.assertEqual(sheet["D10"].value, "ทีม 1")
        self.assertEqual(sheet["I10"].value, "=SUM(E10:H10)")
        self.assertEqual(sheet["AE10"].value, "=I10+N10+S10+X10+AC10")
        self.assertEqual(sheet["AF10"].value, "=RANK(AE10,$AE$10:$AE$10,0)")
        self.assertEqual(sheet["AJ10"].value.text, "=_xlfn.SORTBY(B10:AF10,AF10:AF10,1,C10:C10,1)")
        self.assertEqual(sheet["AM3"].value, "Atelier 1")
        self.assertEqual(sheet["B10"].value, "=AF10")
        active_theme = SiteTheme.query.filter_by(is_active=True).first()
        self.assertEqual(sheet["B10"].fill.fgColor.rgb[-6:], active_theme.color_primary_dark.lstrip("#").upper())
        self.assertGreaterEqual(len(sheet._images), 10)
        self.assertIn(event.name, workbook["TIR REPECHAGE"]["B2"].value)
        self.assertIn(event.name, workbook["TIR FINAL "]["C1"].value)


if __name__ == "__main__":
    unittest.main()
